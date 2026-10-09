"""
td-my-team: fetch_data.py

Downloads the published Google Sheet (.xlsx), parses the Yahoo IMPORTHTML tables,
matches players by name, computes replacement baselines and the two lenses
(FOR ME = team log starts, OVERALL = all games), and writes data/data.json.

Run in GitHub Actions:   python fetch_data.py
Test against a local file: python fetch_data.py --file path/to/workbook.xlsx

No Yahoo API, no cookies. Python never contacts Yahoo. See SPEC.md.
"""

import io
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
from datetime import datetime, timezone, timedelta, time as dtime

import openpyxl

# ----------------------------------------------------------------------------
# CONFIG (see SPEC.md sections 2, 3, 6, 8)
# ----------------------------------------------------------------------------

SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/e/"
    "2PACX-1vQWduNAcz_oG_y6og2MHut-t4mmCTZe8_n6Od2Gzzg7_xHJycjqNlG2ydKUImdaXJpphHxiBvhcESHP"
    "/pub?output=xlsx"
)
OUT_PATH = os.path.join("data", "data.json")

MY_TEAM_NAME = "Soft Dump in Corner"
SEASON_LABEL = "2026-27"

SKATER_WEIGHTS = {"G": 15, "A": 10, "+/-": 5, "PIM": 2.5, "PPP": 5, "SOG": 2}
GOALIE_WEIGHTS = {"W": 20, "GA": -5, "SV": 1, "SHO": 10}

WINDOWS = ["season", "l30", "l14"]
WINDOW_TAB = {"season": "Season", "l30": "L30", "l14": "L14"}
POSITIONS = ["C", "LW", "RW", "D", "G"]
BASE_POSITIONS = ["C", "LW", "RW", "FWD", "D", "G"]
FORWARDS = {"C", "LW", "RW"}

OWN_FLOOR = {"C": 0.20, "LW": 0.20, "RW": 0.20, "FWD": 0.20, "D": 0.33, "G": 0.0}
MIN_GP = {"season": 10, "l30": 3, "l14": 3}
EARLY_SEASON_MIN_GP = 3
POOL_SIZE = 5
SMALL_SAMPLE_STARTS = 5
REPLACE_PTS = -50
TREND_PCT = 0.15
INJURED_TAGS = {"IR", "IR-LT", "IR-NR", "O"}
STATUS_ALIASES = {  # Status Overrides tab values -> site status
    "active": "Active", "injured": "Injured", "ir": "Injured",
    "cut": "Cut", "inactive": "Cut", "dropped": "Cut",
}
STALE_DAYS = 2
MAX_GP = 82

TAB_TEAMLOG = {"skater": "Import My Skaters", "goalie": "Import My Goalies"}
TAB_MYTEAM = {
    "season": "Import My Team Season Avg",
    "l30": "Import My Team L30 Avg",
    "l14": "Import My Team L14 Avg",
}
TAB_OVERRIDES = "Status Overrides"
TAB_NAMEFIX = "Name Fixes"


def fa_tab(pos, window):
    if pos == "G":
        return f"Import FA Goalies-{WINDOW_TAB[window]} Total"
    return f"Import FA {pos}-{WINDOW_TAB[window]} Avg"


TABLE_LABEL = {"skater": "Forwards/Defensemen", "goalie": "Goaltenders"}
REQUIRED_COLS = {
    "skater": ["GP*"] + list(SKATER_WEIGHTS),
    "goalie": ["GP*"] + list(GOALIE_WEIGHTS),
}

# Yahoo team abbreviations (several variants, longest tried first)
TEAMS = {
    "ANA", "ARI", "BOS", "BUF", "CGY", "CAR", "CHI", "COL", "CBJ", "CLS", "DAL",
    "DET", "EDM", "FLA", "LA", "LAK", "MIN", "MTL", "MON", "NJ", "NJD", "NSH",
    "NAS", "NYI", "NYR", "OTT", "PHI", "PIT", "SJ", "SJS", "SEA", "STL", "TB",
    "TBL", "TOR", "UTA", "UTAH", "VAN", "VGK", "VEG", "WSH", "WAS", "WPG",
}
TEAMS_BY_LEN = sorted(TEAMS, key=len, reverse=True)

NOTE_SUFFIXES = ["No new player Notes", "New Player Note", "Player Note"]
TAG_RE = re.compile(r"(IR-LT|IR-NR|IR|DTD|NA|O)$")
POS_RE = re.compile(r" - ((?:C|LW|RW|D|G)(?:,(?:C|LW|RW|D|G))*)\s*$")
PUA_RE = re.compile(r"[\ue000-\uf8ff]")


# ----------------------------------------------------------------------------
# HELPERS
# ----------------------------------------------------------------------------

def clean_text(v):
    if v is None:
        return ""
    return PUA_RE.sub("", str(v)).strip()


def to_num(v):
    """Yahoo cells: numbers, '-', '', or '33%'. Blank or '-' = 0."""
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = clean_text(v).replace(",", "")
    if s in ("", "-", "--", "N/A"):
        return 0.0
    pct = s.endswith("%")
    s = s.rstrip("%")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return x / 100.0 if pct else x


def to_own(v):
    x = to_num(v)
    return x / 100.0 if x > 1.0 else x


def fold(s):
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.lower().replace(".", "").replace("'", "").replace("\u2019", "")
    return re.sub(r"\s+", " ", s).strip()


def pos_group(positions):
    if not positions:
        return "F"
    first = positions[0]
    if first == "G":
        return "G"
    if first == "D":
        return "D"
    return "F"


NAME_FIXES = {}  # folded raw name -> replacement name (filled from Name Fixes tab)


def parse_name(raw):
    """'Sebastian AhoPlayer NoteCAR - C\\n7:00 pm vs VAN' -> dict, or None."""
    if not isinstance(raw, str):
        return None
    line = PUA_RE.sub("", raw.split("\n")[0]).strip()
    m = POS_RE.search(line)
    if not m:
        return None
    positions = m.group(1).split(",")
    head = line[: m.start()]

    team = None
    for ab in TEAMS_BY_LEN:
        if head.upper().endswith(ab) and head[-len(ab):].isalpha():
            team = ab
            break
    if team is None:
        mt = re.search(r"([A-Z]{2,4})$", head)
        if not mt:
            return None
        team = mt.group(1)
    name = head[: -len(team)]

    tag = None
    for _ in range(2):
        for suf in NOTE_SUFFIXES:
            if name.lower().endswith(suf.lower()):
                name = name[: -len(suf)]
        mt = TAG_RE.search(name)
        if mt and tag is None:
            before = name[: mt.start()]
            if before and (before[-1].islower() or before[-1] in ".'"):
                tag = mt.group(1)
                name = before
    name = re.sub(r"\s+", " ", name).strip()
    if not name:
        return None

    fixed = NAME_FIXES.get(fold(name))
    if fixed:
        name = fixed
    group = pos_group(positions)
    return {
        "name": name,
        "team": team,
        "pos": positions,
        "group": group,
        "tag": tag,
        "key": f"{fold(name)}|{group}",
    }


# ----------------------------------------------------------------------------
# WORKBOOK READING
# ----------------------------------------------------------------------------

def download(url, tries=3):
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "td-my-team/1.0"})
            with urllib.request.urlopen(req, timeout=120) as r:
                data = r.read()
            if data[:2] != b"PK":
                raise ValueError("Response is not an .xlsx file (sheet not published?)")
            return data
        except Exception as e:  # noqa: BLE001
            last = e
            print(f"Download attempt {i + 1} failed: {e}")
            time.sleep(5 * (i + 1))
    raise SystemExit(f"ERROR: could not download sheet: {last}")


def sheet_rows(wb, title):
    if title not in wb.sheetnames:
        return None
    return [list(r) for r in wb[title].iter_rows(values_only=True)]


def cell(rows, r, c):
    if r < 0 or r >= len(rows):
        return None
    row = rows[r]
    return row[c] if 0 <= c < len(row) else None


def find_table(rows, kind):
    """Find the first complete Yahoo table of this kind.

    Avg/total layout: label cell and column names on the same row ('GP*' nearby).
    Team log layout:  label row, then a names row with 'Name' one column right.
    Helper columns ('3 lines', 'Clean 1', ...) and mirrored copies are skipped
    because they lack the required stat columns.
    Returns (records, column_map, pull_date_text) or (None, None, None).
    """
    if not rows:
        return None, None, None
    label = TABLE_LABEL[kind]
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if clean_text(v) != label:
                continue
            same_row = [clean_text(cell(rows, r, j)) for j in range(c, c + 7)]
            if "GP*" in same_row:
                hdr, namecol = r, c
            elif clean_text(cell(rows, r + 1, c + 1)) == "Name":
                hdr, namecol = r + 1, c + 1
            else:
                continue
            cmap, opp = {}, None
            for j in range(c, c + 25):
                h = clean_text(cell(rows, hdr, j))
                if not h:
                    continue
                if h in ("3 lines", "2 lines", "Prelim") or h.startswith("Clean"):
                    break
                if h.startswith("Opp:"):
                    opp = h
                if h in cmap:
                    break
                cmap[h] = j
            if any(k not in cmap for k in REQUIRED_COLS[kind]):
                continue
            recs, started = [], False
            for i in range(hdr + 1, len(rows)):
                raw = cell(rows, i, namecol)
                if isinstance(raw, str) and " - " in raw:
                    started = True
                    rec = {"raw": raw}
                    for h, j in cmap.items():
                        rec[h] = cell(rows, i, j)
                    recs.append(rec)
                elif started:
                    break
                elif i > hdr + 3:
                    break
            return recs, cmap, opp
    return None, None, None


def fan_pts_from_stats(rec, kind):
    w = SKATER_WEIGHTS if kind == "skater" else GOALIE_WEIGHTS
    return sum(to_num(rec.get(k)) * wt for k, wt in w.items())


def toi_minutes(v):
    """Yahoo 'mm:ss' arrives from Google as a time or duration read as h:mm.
    Returns minutes, or None if blank."""
    if v is None or v == "":
        return None
    if isinstance(v, timedelta):
        return v.total_seconds() / 3600.0
    if isinstance(v, dtime):
        return v.hour + v.minute / 60.0 + v.second / 3600.0
    if isinstance(v, datetime):
        return v.hour + v.minute / 60.0 + v.second / 3600.0
    if isinstance(v, (int, float)):
        return float(v) * 24.0 if float(v) < 5 else float(v)
    m = re.match(r"^\s*(\d+):(\d{1,2})", str(v))
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60.0
    return None


def stat_line(rec, kind, per_game_source, gp):
    """Per-game stat line. per_game_source=True when the tab already holds averages.
    Goalie TOI* is always a season total in Yahoo's tables."""
    if gp <= 0:
        return None
    keys = SKATER_WEIGHTS if kind == "skater" else GOALIE_WEIGHTS
    div = 1.0 if per_game_source else gp
    line = {k: to_num(rec.get(k)) / div for k in keys}
    if kind == "skater":
        toi = toi_minutes(rec.get("TOI/G*"))
        line["TOI"] = toi
    else:
        ga_tot, sv_tot = line["GA"] * gp, line["SV"] * gp
        line["SVP"] = sv_tot / (sv_tot + ga_tot) if (sv_tot + ga_tot) > 0 else None
        toi = toi_minutes(rec.get("TOI*"))
        line["GAA"] = (ga_tot * 60.0 / toi) if toi else None
    return line


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------

def main():
    args = sys.argv[1:]
    if "--file" in args:
        path = args[args.index("--file") + 1]
        with open(path, "rb") as f:
            blob = f.read()
        source = os.path.basename(path)
    else:
        blob = download(SHEET_URL)
        source = "published Google Sheet"

    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    health = {"tabs": [], "warnings": [], "unparsed": [], "name_checks": []}
    pull_dates = set()

    def note_tab(title, kind, n, required=True):
        health["tabs"].append({"tab": title, "kind": kind, "rows": n})
        if n is None:
            if required:
                health["warnings"].append(f"Missing tab: {title}")
        elif n == 0:
            health["warnings"].append(f"Empty table: {title} ({kind})")

    def parsed(recs, title):
        out = []
        for rec in recs or []:
            p = parse_name(rec["raw"])
            if p is None:
                health["unparsed"].append({"tab": title, "raw": clean_text(rec["raw"])[:80]})
                continue
            p["rec"] = rec
            out.append(p)
        return out

    # ---- optional tabs -----------------------------------------------------
    overrides = {}
    rows = sheet_rows(wb, TAB_NAMEFIX)
    if rows:
        for r in rows[1:]:
            if len(r) >= 2 and r[0] and r[1]:
                NAME_FIXES[fold(clean_text(r[0]))] = clean_text(r[1])
    rows = sheet_rows(wb, TAB_OVERRIDES)
    if rows:
        for r in rows[1:]:
            if len(r) >= 2 and r[0] and r[1]:
                overrides[fold(clean_text(r[0]))] = clean_text(r[1]).title()

    # ---- my team (OVERALL) -------------------------------------------------
    players = {}

    def get_player(p):
        pl = players.get(p["key"])
        if pl is None:
            pl = {
                "key": p["key"], "name": p["name"], "team": p["team"],
                "pos": p["pos"], "group": p["group"], "tag": p["tag"],
                "on_roster": False, "overall": {}, "forme": None,
            }
            players[p["key"]] = pl
        return pl

    scoring_mismatch = 0
    for w in WINDOWS:
        title = TAB_MYTEAM[w]
        rows = sheet_rows(wb, title)
        for kind in ("skater", "goalie"):
            recs, cmap, opp = find_table(rows, kind)
            note_tab(title, kind, None if rows is None else len(recs or []))
            if opp:
                pull_dates.add(opp)
            for p in parsed(recs, title):
                rec = p["rec"]
                gp = to_num(rec.get("GP*"))
                ppg = to_num(rec.get("Fan Pts"))
                if gp > 0 and abs(fan_pts_from_stats(rec, kind) - ppg) > 0.5:
                    scoring_mismatch += 1
                if w == "season" and gp > MAX_GP:
                    health["warnings"].append(
                        f"{title}: {p['name']} shows {int(gp)} GP. Season link may point to last season.")
                pl = get_player(p)
                pl["on_roster"] = True
                # roster row is the freshest source for team / tag / positions
                pl.update({"team": p["team"], "pos": p["pos"], "tag": p["tag"]})
                pl["overall"][w] = {
                    "gp": gp,
                    "ppg": ppg if gp > 0 else None,
                    "own": to_own(rec.get("% Ros")),
                    "stats": stat_line(rec, kind, True, gp),
                }
    if scoring_mismatch:
        health["warnings"].append(
            f"{scoring_mismatch} my-team rows where Yahoo Fan Pts differ from the scoring weights. "
            "Check league scoring settings.")

    # ---- team log (FOR ME) -------------------------------------------------
    for kind, title in TAB_TEAMLOG.items():
        rows = sheet_rows(wb, title)
        recs, _, _ = find_table(rows, kind)
        note_tab(title, kind, None if rows is None else len(recs or []))
        for p in parsed(recs, title):
            rec = p["rec"]
            gp = to_num(rec.get("GP*"))
            pts = fan_pts_from_stats(rec, kind)
            pl = get_player(p)
            pl["forme"] = {"gp": gp, "pts": pts, "ppg": pts / gp if gp > 0 else None,
                           "stats": stat_line(rec, kind, False, gp)}

    roster_count = sum(1 for p in players.values() if p["on_roster"])
    if roster_count == 0:
        raise SystemExit("ERROR: no players found on the my-team tabs. Not overwriting the site.")

    # ---- free agents + baselines ------------------------------------------
    fa = {w: {} for w in WINDOWS}  # fa[w][pos] = list of rows
    for w in WINDOWS:
        for pos in POSITIONS:
            title = fa_tab(pos, w)
            kind = "goalie" if pos == "G" else "skater"
            rows = sheet_rows(wb, title)
            recs, cmap, opp = find_table(rows, kind)
            note_tab(title, kind, None if rows is None else len(recs or []))
            if opp:
                pull_dates.add(opp)
            out = []
            for p in parsed(recs, title):
                rec = p["rec"]
                gp = to_num(rec.get("GP*"))
                fp = to_num(rec.get("Fan Pts"))
                if pos == "G":          # goalie tabs are totals
                    total, ppg = fp, (fp / gp if gp > 0 else 0.0)
                else:                   # skater tabs are per-game averages
                    ppg, total = fp, fp * gp
                if w == "season" and gp > MAX_GP:
                    health["warnings"].append(
                        f"{title}: {p['name']} shows {int(gp)} GP. Season link may point to last season.")
                out.append({
                    "key": p["key"], "name": p["name"], "team": p["team"], "pos": p["pos"],
                    "tag": p["tag"], "own": to_own(rec.get("% Ros")),
                    "gp": gp, "ppg": ppg, "total": total,
                    "stats": stat_line(rec, kind, pos != "G", gp),
                })
            fa[w][pos] = out

    baselines = {w: {} for w in WINDOWS}
    for w in WINDOWS:
        for bpos in BASE_POSITIONS:
            if bpos == "FWD":
                seen, rows_ = set(), []
                for p in ("C", "LW", "RW"):
                    for r in fa[w].get(p, []):
                        if r["key"] not in seen:
                            seen.add(r["key"])
                            rows_.append(r)
            else:
                rows_ = fa[w].get(bpos, [])
            floor = OWN_FLOOR[bpos]
            min_gp, early = MIN_GP[w], False
            qual = [r for r in rows_ if r["own"] >= floor and r["gp"] >= min_gp]
            if w == "season" and len(qual) < POOL_SIZE:
                min_gp, early = EARLY_SEASON_MIN_GP, True
                qual = [r for r in rows_ if r["own"] >= floor and r["gp"] >= min_gp]
            pool = sorted(qual, key=lambda r: r["ppg"], reverse=True)[:POOL_SIZE]
            gp_sum = sum(r["gp"] for r in pool)
            baselines[w][bpos] = {
                "ppg": (sum(r["total"] for r in pool) / gp_sum) if gp_sum else None,
                "n": len(pool),
                "min_gp": min_gp,
                "floor": floor,
                "early": early,
                "members": [r["key"] for r in pool],
            }

    # ---- comparisons (per eligible position + all forwards) ---------------
    def compare(pl, w, ppg, gp):
        """Returns {pos: {base, gap, pts}} for every eligible position with a line,
        plus 'FWD' for forwards. The site picks the one matching the row it shows."""
        if ppg is None or not gp:
            return {}
        out = {}
        keys = [p for p in pl["pos"] if p in POSITIONS]
        if pl["group"] == "F":
            keys.append("FWD")
        for k in keys:
            base = baselines[w].get(k, {}).get("ppg")
            if base is not None:
                out[k] = {"base": base, "gap": ppg - base, "pts": (ppg - base) * gp}
        return out

    out_players = []
    for pl in players.values():
        fkey = fold(pl["name"])
        if fkey in overrides:
            status = overrides[fkey]
        elif pl["on_roster"]:
            status = "Injured" if pl["tag"] in INJURED_TAGS else "Active"
        else:
            status = "Cut"
        status = STATUS_ALIASES.get(status.lower(), "Active")

        forme = None
        if pl["forme"] and pl["forme"]["gp"] > 0:
            f = dict(pl["forme"])
            f["cmp"] = compare(pl, "season", f["ppg"], f["gp"])
            forme = f

        overall = None
        if pl["on_roster"]:
            overall = {}
            for w in WINDOWS:
                o = pl["overall"].get(w)
                if not o:
                    overall[w] = None
                    continue
                o = dict(o)
                o["cmp"] = compare(pl, w, o["ppg"], o["gp"])
                overall[w] = o

        # flags per comparison key (row-level: Tkachuk can be flagged at LW, not RW)
        flags_by = {}
        l30 = overall.get("l30") if overall else None
        if l30:
            for k, c in l30["cmp"].items():
                fl = []
                if c["gap"] < 0:
                    fl.append("poor")
                if c["pts"] < REPLACE_PTS:
                    fl.append("replace")
                flags_by[k] = fl
        trend = None
        if overall:
            ssn = overall.get("season")
            if l30 and ssn and l30.get("ppg") is not None and ssn.get("ppg"):
                trend = l30["ppg"] / ssn["ppg"] - 1
        pflags = []
        if forme and forme["gp"] < SMALL_SAMPLE_STARTS:
            pflags.append("small")
        if not forme and pl["on_roster"]:
            pflags.append("nostarts")

        out_players.append({
            "key": pl["key"], "name": pl["name"], "team": pl["team"], "pos": pl["pos"],
            "group": pl["group"], "tag": pl["tag"], "status": status,
            "forme": forme, "overall": overall, "trend": trend,
            "flags_by": flags_by, "pflags": pflags,
        })

    # possible name mismatches: a Cut player sharing a last name + group with a roster player
    roster_last = {}
    for p in out_players:
        if p["status"] != "Cut":
            roster_last.setdefault((fold(p["name"]).split(" ")[-1], p["group"]), []).append(p["name"])
    for p in out_players:
        if p["status"] == "Cut":
            hits = roster_last.get((fold(p["name"]).split(" ")[-1], p["group"]))
            if hits:
                health["name_checks"].append(
                    f"Team log '{p['name']}' may be the same player as roster '{hits[0]}'. "
                    "If so, add a row to a 'Name Fixes' tab.")

    # ---- waiver wire (C, LW, RW, F, D, G) ---------------------------------
    active = [p for p in out_players if p["status"] == "Active" and p["overall"]]
    waiver = {w: {} for w in WINDOWS}
    for w in WINDOWS:
        for wpos in POSITIONS + ["F"]:
            bkey = "FWD" if wpos == "F" else wpos
            b = baselines[w][bkey]
            if wpos == "F":
                seen, src = set(), []
                for p in ("C", "LW", "RW"):
                    for r in fa[w].get(p, []):
                        if r["key"] not in seen:
                            seen.add(r["key"])
                            src.append(r)
                elig = lambda p: p["group"] == "F"
            else:
                src = fa[w].get(wpos, [])
                elig = lambda p, wp=wpos: wp in p["pos"]
            mine = [p for p in active if elig(p)
                    and p["overall"].get(w) and p["overall"][w]["ppg"] is not None]
            weakest = min(mine, key=lambda p: p["overall"][w]["ppg"]) if mine else None
            rows_ = []
            for r in sorted(src, key=lambda r: r["ppg"], reverse=True):
                reasons = []
                if r["own"] < b["floor"]:
                    reasons.append(f"under {int(round(b['floor'] * 100))}% owned")
                if r["gp"] < b["min_gp"]:
                    reasons.append(f"under {b['min_gp']} GP")
                upgrade = None
                if not reasons and weakest:
                    gain = r["ppg"] - weakest["overall"][w]["ppg"]
                    if gain > 0:
                        upgrade = {"over": weakest["name"], "gain": gain}
                rows_.append({
                    **r,
                    "gap": (r["ppg"] - b["ppg"]) if b["ppg"] is not None else None,
                    "pool": r["key"] in b["members"],
                    "excluded": reasons,
                    "upgrade": upgrade,
                })
            waiver[w][wpos] = rows_

    # ---- health ------------------------------------------------------------
    now_utc = datetime.now(timezone.utc)
    now_et = now_utc - timedelta(hours=4)  # display only; DST drift of 1h is harmless
    pull = None
    if pull_dates:
        txt = sorted(pull_dates)[0]
        m = re.search(r"(\d{1,2})/(\d{1,2})", txt)
        if m:
            mo, dy = int(m.group(1)), int(m.group(2))
            yr = now_et.year - (1 if mo > now_et.month + 1 else 0)
            pull = datetime(yr, mo, dy).date()
        if len(pull_dates) > 1:
            health["warnings"].append(
                f"Tabs show different pull dates ({', '.join(sorted(pull_dates))}). Some imports may be stale.")
    if pull is None:
        health["warnings"].append("Could not read the pull date from the 'Opp:' header.")
    elif (now_et.date() - pull).days > STALE_DAYS:
        health["warnings"].append(
            f"Sheet data is from {pull.isoformat()}, more than {STALE_DAYS} days old. Open the sheet to refresh imports.")

    data = {
        "generated_utc": now_utc.isoformat(timespec="seconds"),
        "generated_et": now_et.strftime("%Y-%m-%d %H:%M"),
        "pull_date": pull.isoformat() if pull else None,
        "source": source,
        "team_name": MY_TEAM_NAME,
        "season": SEASON_LABEL,
        "config": {
            "own_floor": OWN_FLOOR, "min_gp": MIN_GP, "early_min_gp": EARLY_SEASON_MIN_GP,
            "pool_size": POOL_SIZE, "small_sample": SMALL_SAMPLE_STARTS,
            "replace_pts": REPLACE_PTS, "trend_pct": TREND_PCT,
            "skater_weights": SKATER_WEIGHTS, "goalie_weights": GOALIE_WEIGHTS,
        },
        "baselines": baselines,
        "players": out_players,
        "waiver": waiver,
        "health": health,
    }
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)

    print(f"Source: {source}")
    print(f"Players: {len(out_players)} ({roster_count} on roster)")
    for w in WINDOWS:
        line = ", ".join(
            f"{p} {baselines[w][p]['ppg']:.1f}" if baselines[w][p]["ppg"] is not None else f"{p} n/a"
            for p in BASE_POSITIONS)
        print(f"Baselines {w}: {line}")
    for msg in health["warnings"] + health["name_checks"]:
        print("WARNING:", msg)
    for u in health["unparsed"]:
        print("UNPARSED:", u)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
